"""Export the report's Markdown to PDF and DOCX.

    uv run --frozen --group report --with playwright==1.64.1 python tools/build_report.py

Input: ``docs/report/earnings-vol-report.md`` (rendered from its template by ``tools/report_numbers.py``; run that
first, or ``--check`` it). Outputs, next to it:

- ``earnings-vol-report.pdf``: the Markdown is turned into HTML with markdown-it-py (CommonMark plus tables), styled
  for print (``CSS``), and printed to A4 by headless Chromium through Playwright 1.64.1 (``page.pdf``). The browser
  binary is ``$PLAYWRIGHT_CHROMIUM`` or ``~/.cache/ms-playwright/chromium-1234/chrome-linux/chrome``. The page loads
  only local files: the HTML and the PNGs in ``docs/images/``.
- ``earnings-vol-report.docx``: the same markdown-it token stream written with python-docx (``DocxWriter``):
  headings, paragraphs with bold, italic, code and links, bullet and numbered lists, block quotes, tables, code
  blocks and images. HTML comments are dropped.

pandoc is not installed on the build machine, so it is not used. Each build records, in
``earnings-vol-report.build.json``, the SHA-256 of the export and of the Markdown it was built from;
``tools/report_numbers.py --check`` fails when an export is missing, was changed after the build, or was built from
an older Markdown. Both files must stay under ``MAX_BYTES`` (10 MB);
the script fails otherwise. The PDF and DOCX embed a creation time, so they are not byte-for-byte reproducible;
the Markdown is (``tools/report_numbers.py --check``).

Dependencies: the ``report`` group of ``pyproject.toml`` (python-docx, markdown-it-py), and Playwright added for the
run with ``--with playwright==1.64.1``.
"""

from __future__ import annotations

import argparse
import hashlib
import html
import json
import os
import re
import sys
import tempfile
from pathlib import Path

REPO = Path(__file__).resolve().parents[1]
REPORT_MD = REPO / "docs/report/earnings-vol-report.md"
STAMP = REPO / "docs/report/earnings-vol-report.build.json"  # which Markdown each export was built from
DEFAULT_CHROME = Path.home() / ".cache/ms-playwright/chromium-1234/chrome-linux/chrome"
MAX_BYTES = 10 * 1024 * 1024
TITLE = "Earnings and implied volatility: an event study"
FOOTER = "Draft for owner review (2026-10-29) · not investment advice · results provisional"
# A relative link in the PDF would resolve to this machine's file:// path: point it at the public repository.
PUBLIC_BLOB = "https://github.com/rajib2k5/earnings-vol-study/blob/main/"

CSS = """
@page { size: A4; margin: 18mm 17mm 20mm 17mm; }
html { font-family: "DejaVu Sans", "Liberation Sans", Arial, sans-serif; font-size: 10pt; color: #1b1f24; }
body { margin: 0; background: #ffffff; line-height: 1.42; }
h1 { font-size: 19pt; margin: 6pt 0 8pt; }
h2 { font-size: 14pt; margin: 16pt 0 6pt; border-bottom: 1px solid #c9d1d9; padding-bottom: 2pt; break-after: avoid; }
h3 { font-size: 11.5pt; margin: 12pt 0 4pt; break-after: avoid; }
p, li { orphans: 3; widows: 3; }
blockquote { margin: 8pt 0; padding: 4pt 10pt; border-left: 3px solid #b8860b; background: #fbf6ea; }
code { font-family: "DejaVu Sans Mono", "Liberation Mono", monospace; font-size: 8.6pt; background: #f2f4f7;
       padding: 0 2px; border-radius: 2px; }
pre { background: #f2f4f7; padding: 6pt 8pt; white-space: pre-wrap; word-break: break-word; font-size: 8pt;
      break-inside: avoid; }
pre code { background: none; padding: 0; font-size: 8pt; }
table { border-collapse: collapse; margin: 6pt 0 10pt; font-size: 8.6pt; width: 100%; break-inside: auto; }
th, td { border: 1px solid #c9d1d9; padding: 2.5pt 4pt; vertical-align: top; }
th { background: #eef1f5; }
tr { break-inside: avoid; }
img { display: block; max-width: 100%; max-height: 120mm; margin: 6pt auto 2pt; object-fit: contain;
      break-inside: avoid; }
p:has(> img) { text-align: center; font-size: 8.5pt; color: #57606a; break-inside: avoid; }
hr { border: 0; border-top: 1px solid #d0d7de; margin: 10pt 0; }
a { color: #0b5cad; text-decoration: none; }
"""


def md_parser():
    from markdown_it import MarkdownIt

    return MarkdownIt("commonmark", {"html": True}).enable("table")


def public_href(href: str) -> str:
    """A link target as the PDF should carry it: a relative path becomes its page in the public repository."""
    if re.match(r"[a-z][a-z0-9+.-]*:|#", href):
        return href
    rel = os.path.normpath(os.path.join(REPORT_MD.parent.relative_to(REPO).as_posix(), href)).replace(os.sep, "/")
    return PUBLIC_BLOB + rel


def to_html(md_text: str) -> str:
    body = md_parser().render(md_text)
    body = re.sub(r'<a href="([^"]+)"', lambda m: f'<a href="{html.escape(public_href(html.unescape(m.group(1))))}"',
                  body)
    # An image's alt text doubles as its caption in print: put it under the picture.
    body = re.sub(r'(<img src="[^"]+" alt="([^"]*)"[^>]*>)', lambda m: f"{m.group(1)}<br>{m.group(2)}", body)
    return (f'<!doctype html><html lang="en"><head><meta charset="utf-8"><title>{html.escape(TITLE)}</title>'
            f"<style>{CSS}</style></head><body>{body}</body></html>")


def build_pdf(md_text: str, out: Path, chrome: Path) -> None:
    from playwright.sync_api import sync_playwright

    page_html = to_html(md_text)
    # The HTML is written next to the Markdown, so its relative image paths (../images/...) resolve.
    with tempfile.NamedTemporaryFile("w", suffix=".html", dir=REPORT_MD.parent, delete=False, encoding="utf-8") as f:
        f.write(page_html)
        tmp = Path(f.name)
    try:
        with sync_playwright() as pw:
            browser = pw.chromium.launch(executable_path=str(chrome))
            page = browser.new_page()
            failed: list[str] = []
            page.on("requestfailed", lambda r: failed.append(r.url))
            page.goto(tmp.as_uri(), wait_until="load")
            page.wait_for_function("() => [...document.images].every((i) => i.complete && i.naturalWidth > 0)")
            if failed:
                raise RuntimeError(f"the report page could not load: {failed}")
            page.pdf(
                path=str(out), format="A4", print_background=True, display_header_footer=True,
                margin={"top": "18mm", "bottom": "20mm", "left": "17mm", "right": "17mm"},
                header_template="<span></span>",
                footer_template=(
                    '<div style="width:100%;font-size:7.5px;color:#57606a;text-align:center;font-family:sans-serif">'
                    f'{html.escape(FOOTER)} · page <span class="pageNumber"></span> of '
                    '<span class="totalPages"></span></div>'),
            )
            browser.close()
    finally:
        tmp.unlink(missing_ok=True)


class DocxWriter:
    """Writes markdown-it tokens into a python-docx Document. Only the constructs the report uses are handled;
    anything else raises, so a new construct in the template cannot vanish silently."""

    def __init__(self, base: Path):
        from docx import Document
        from docx.enum.section import WD_ORIENT
        from docx.shared import Cm, Pt

        self.base = base
        self.doc = Document()
        sec = self.doc.sections[0]
        sec.orientation = WD_ORIENT.PORTRAIT
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        for side in ("left_margin", "right_margin"):
            setattr(sec, side, Cm(1.8))
        sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(2.0)
        normal = self.doc.styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(10)
        self.doc.core_properties.title = TITLE
        self.doc.core_properties.author = "earnings-vol-study"
        self.doc.core_properties.comments = FOOTER
        footer = sec.footer.paragraphs[0]
        footer.text = FOOTER
        footer.runs[0].font.size = Pt(7.5)
        self.text_width_cm = 21.0 - 2 * 1.8
        self.list_stack: list[dict] = []
        self.quote = 0
        self.para = None

    # inline -------------------------------------------------------------------------------------------------
    def inline(self, para, children) -> None:
        from docx.shared import Pt, RGBColor

        bold = italic = False
        link: str | None = None
        for t in children:
            if t.type == "text":
                run = para.add_run(t.content)
                run.bold, run.italic = bold or None, italic or None
                if link:
                    run.font.color.rgb = RGBColor(0x0B, 0x5C, 0xAD)
            elif t.type == "code_inline":
                run = para.add_run(t.content)
                run.font.name = "Consolas"
                run.font.size = Pt(8.5)
                run.bold = bold or None
            elif t.type in ("softbreak",):
                para.add_run(" ")
            elif t.type == "hardbreak":
                para.add_run().add_break()
            elif t.type == "strong_open":
                bold = True
            elif t.type == "strong_close":
                bold = False
            elif t.type == "em_open":
                italic = True
            elif t.type == "em_close":
                italic = False
            elif t.type == "link_open":
                link = t.attrs.get("href")
            elif t.type == "link_close":
                link = None
            elif t.type == "image":
                self.image(t)
            elif t.type == "html_inline":
                continue
            else:
                raise ValueError(f"inline token not handled: {t.type}")

    def image(self, tok) -> None:
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Cm, Pt
        from PIL import Image

        path = (self.base / tok.attrs["src"]).resolve()
        with Image.open(path) as im:
            w, h = im.size
        max_w, max_h = self.text_width_cm, 12.0
        width = min(max_w, max_h * w / h)
        pic = self.doc.add_paragraph()
        pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pic.add_run().add_picture(str(path), width=Cm(width))
        alt = "".join(c.content for c in (tok.children or []) if c.type == "text")
        if alt:
            cap = self.doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = cap.add_run(alt)
            run.italic = True
            run.font.size = Pt(8.5)

    # blocks -------------------------------------------------------------------------------------------------
    def new_paragraph(self):
        from docx.shared import Cm

        if self.list_stack:
            lst = self.list_stack[-1]
            depth = len(self.list_stack)
            if lst["ordered"]:
                p = self.doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.7 * depth)
                p.paragraph_format.first_line_indent = Cm(-0.5)
                if lst["fresh"]:
                    p.add_run(f"{lst['n']}. ")
            else:
                style = "List Bullet" if depth == 1 else "List Bullet 2"
                p = self.doc.add_paragraph(style=style) if lst["fresh"] else self.doc.add_paragraph()
                if not lst["fresh"]:
                    p.paragraph_format.left_indent = Cm(0.63 * depth)
            lst["fresh"] = False
            return p
        if self.quote:
            p = self.doc.add_paragraph(style="Quote")
            return p
        return self.doc.add_paragraph()

    def table(self, tokens, i: int) -> int:
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt

        rows: list[list[tuple[list, str]]] = []
        header_rows = 0
        in_head = False
        j = i + 1
        while tokens[j].type != "table_close":
            t = tokens[j]
            if t.type == "thead_open":
                in_head = True
            elif t.type == "thead_close":
                in_head = False
            elif t.type == "tr_open":
                rows.append([])
                header_rows += in_head
            elif t.type in ("th_open", "td_open"):
                style = t.attrs.get("style") or ""
                align = "right" if "right" in style else ("center" if "center" in style else "left")
                rows[-1].append((tokens[j + 1].children or [], align))
            j += 1
        ncol = max(len(r) for r in rows)
        tbl = self.doc.add_table(rows=len(rows), cols=ncol)
        tbl.style = "Table Grid"
        for r, row in enumerate(rows):
            for c, (children, align) in enumerate(row):
                cell = tbl.cell(r, c)
                para = cell.paragraphs[0]
                para.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "right": WD_ALIGN_PARAGRAPH.RIGHT,
                                  "center": WD_ALIGN_PARAGRAPH.CENTER}[align]
                self.inline(para, children)
                for run in para.runs:
                    run.font.size = Pt(8)
                    if r < header_rows:
                        run.bold = True
        self.doc.add_paragraph()
        return j

    def code_block(self, content: str) -> None:
        from docx.shared import Pt

        for line in content.rstrip("\n").split("\n"):
            p = self.doc.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(line)
            run.font.name = "Consolas"
            run.font.size = Pt(7.5)
        self.doc.add_paragraph()

    def write(self, tokens) -> None:
        i = 0
        while i < len(tokens):
            t = tokens[i]
            if t.type == "heading_open":
                level = int(t.tag[1])
                h = self.doc.add_heading(level=min(level - 1, 3) if level > 1 else 0)
                self.inline(h, tokens[i + 1].children or [])
                i += 2
            elif t.type == "paragraph_open":
                inl = tokens[i + 1].children or []
                if len(inl) == 1 and inl[0].type == "image":
                    self.image(inl[0])
                else:
                    self.inline(self.new_paragraph(), inl)
                i += 2
            elif t.type in ("bullet_list_open", "ordered_list_open"):
                start = int(t.attrs.get("start", 1)) if t.type == "ordered_list_open" else 1
                self.list_stack.append({"ordered": t.type == "ordered_list_open", "n": start - 1, "fresh": False})
            elif t.type in ("bullet_list_close", "ordered_list_close"):
                self.list_stack.pop()
            elif t.type == "list_item_open":
                self.list_stack[-1]["n"] += 1
                self.list_stack[-1]["fresh"] = True
            elif t.type == "blockquote_open":
                self.quote += 1
            elif t.type == "blockquote_close":
                self.quote -= 1
            elif t.type == "table_open":
                i = self.table(tokens, i)
            elif t.type in ("fence", "code_block"):
                self.code_block(t.content)
            elif t.type == "hr":
                self.doc.add_paragraph()
            elif t.type in ("html_block", "list_item_close", "paragraph_close", "heading_close", "inline"):
                pass
            else:
                raise ValueError(f"block token not handled: {t.type}")
            i += 1

    def save(self, out: Path) -> None:
        self.doc.save(str(out))


def build_docx(md_text: str, out: Path) -> None:
    w = DocxWriter(REPORT_MD.parent)
    w.write(md_parser().parse(md_text))
    w.save(out)


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--only", choices=("pdf", "docx"), help="build one format")
    args = ap.parse_args(argv)
    md_text = REPORT_MD.read_text(encoding="utf-8")
    if "{{" in md_text:
        print("FAIL: the report still has a placeholder; run tools/report_numbers.py first", file=sys.stderr)
        return 1
    outs = []
    if args.only in (None, "docx"):
        out = REPORT_MD.with_suffix(".docx")
        build_docx(md_text, out)
        outs.append(out)
    if args.only in (None, "pdf"):
        chrome = Path(os.environ.get("PLAYWRIGHT_CHROMIUM", DEFAULT_CHROME))
        if not chrome.exists():
            print(f"FAIL: no Chromium at {chrome}; set PLAYWRIGHT_CHROMIUM", file=sys.stderr)
            return 1
        out = REPORT_MD.with_suffix(".pdf")
        build_pdf(md_text, out, chrome)
        outs.append(out)
    stamp = json.loads(STAMP.read_text()) if STAMP.exists() else {}
    md_sha = hashlib.sha256(md_text.encode("utf-8")).hexdigest()
    for out in outs:
        stamp[out.suffix.lstrip(".")] = {"file": out.name, "sha256": hashlib.sha256(out.read_bytes()).hexdigest(),
                                         "markdown_sha256": md_sha}
    STAMP.write_text(json.dumps(dict(sorted(stamp.items())), indent=2) + "\n")
    for out in outs:
        size = out.stat().st_size
        print(f"wrote {out.relative_to(REPO)} ({size:,} B)")
        if size >= MAX_BYTES:
            print(f"FAIL: {out.name} is {size:,} B, over {MAX_BYTES:,}", file=sys.stderr)
            return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
