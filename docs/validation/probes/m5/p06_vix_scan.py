"""p06: does any published text carry a raw VIX value or a per-event VIX-adjusted value?

    uv run --frozen --offline --no-sync python docs/validation/probes/m5/p06_vix_scan.py [--plant]

Reference sets (local stage cache, never printed): every VIX close and ln VIX (market_daily), and every per-event
*_cboe value of event_measures (the VIX-adjusted crush, run-up, persistence, dsigma, beta, standardised values),
each formatted the ways the docs print numbers (2, 3 and 4 decimals, with an optional sign and a Unicode minus).
Texts scanned: the report Markdown, README, docs/prfaq.md, the PDF text (pdftotext) and the DOCX text.
Each number token that hits a reference set is classified: "generated" if the same string is the rendered value of
a placeholder whose source is an aggregate (from p01's output), otherwise "unexplained" (printed with file:line, not
the value). --plant: one per-event VIX-adjusted crush value is planted into a scratch copy of the report text; the
scan must then report one unexplained hit (the falsifier).
"""
from __future__ import annotations

import re
import subprocess
import sys
from pathlib import Path

import pandas as pd

REPO = Path(__file__).resolve().parents[4]
ST = Path("/home/rajib2k5/data/vol-projects/earnings-vol-study/cache/stage/74bed6f4b1281810")
M = "−"


def forms(x: float) -> set[str]:
    out = set()
    for nd in (2, 3, 4):
        s = f"{x:.{nd}f}"
        out.add(s.replace("-", M))
        out.add(s.lstrip("-"))  # an unsigned print of a negative value
        if not s.startswith("-"):
            out.add("+" + s)
    return out


def main() -> int:
    plant = "--plant" in sys.argv
    md = pd.read_parquet(ST / "market_daily.parquet")
    vix = set()
    for c in ("vix_close_cboe", "ln_vix_cboe"):
        for v in md[c].dropna():
            vix |= forms(float(v))
    em = pd.read_parquet(ST / "event_measures.parquet")
    cols = [c for c in em.columns if c.endswith("_cboe")]
    per = set()
    for c in cols:
        for v in em[c].dropna():
            per |= forms(float(v))
    gen_str = re.findall(r"got='([^']*)'", (REPO / "docs/validation/probes/m5/out/p01_oracle.txt").read_text())
    # every number inside a generated string (interval ends included), signs dropped
    gen = {t.lstrip("+" + M + "-") for g in gen_str for t in re.findall(r"[−+-]?\d+\.\d+", g)}
    strict = "--strict" in sys.argv
    texts = {
        "report.md": (REPO / "docs/report/earnings-vol-report.md").read_text(),
        "README.md": (REPO / "README.md").read_text(),
        "prfaq.md": (REPO / "docs/prfaq.md").read_text(),
        "report.pdf": subprocess.run(["pdftotext", "-layout", str(REPO / "docs/report/earnings-vol-report.pdf"), "-"],
                                     capture_output=True, text=True).stdout,
    }
    import docx
    d = docx.Document(str(REPO / "docs/report/earnings-vol-report.docx"))
    texts["report.docx"] = "\n".join([p.text for p in d.paragraphs] +
                                     [c.text for t in d.tables for r in t.rows for c in r.cells])
    if plant:
        v = float(em.loc[em["crush_cboe"].notna(), "crush_cboe"].iloc[17])
        texts["report.md"] += f"\nplanted: the AAPL crush was {f'{v:.4f}'.replace('-', M)} log points.\n"
    print(f"reference forms: VIX {len(vix):,}; per-event VIX-adjusted ({len(cols)} columns) {len(per):,}")
    total_unexpl = 0
    for name, t in texts.items():
        hits_g = hits_u = 0
        where = []
        for i, line in enumerate(t.splitlines(), 1):
            for tok, pc in re.findall(r"([−+-]?\d+\.\d+)(%?)", line):
                tok = tok.replace("-", M) if tok.startswith("-") else tok
                dec = len(tok.split(".")[1])
                if strict:  # discriminating forms only: a per-event value printed to 4+ decimals, or a
                    # VIX-like level (two decimals, 9 to 90, not a percentage)
                    is_vix = dec == 2 and not pc and 9 <= float(tok.lstrip("+" + M)) <= 90 and tok in vix
                    is_per = dec >= 4 and tok in per
                else:
                    is_vix, is_per = tok in vix, tok in per
                if is_vix or is_per:
                    if tok.lstrip("+" + M) in gen:
                        hits_g += 1
                    else:
                        hits_u += 1
                        where.append(f"{name}:{i} ({'VIX' if is_vix else 'per-event'} form)")
        total_unexpl += hits_u
        print(f"{name:12} hits that are a generated aggregate's string: {hits_g:3}; unexplained hits: {hits_u}"
              + (f" at {where[:12]}" if where else ""))
    print(f"UNEXPLAINED TOTAL: {total_unexpl}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
